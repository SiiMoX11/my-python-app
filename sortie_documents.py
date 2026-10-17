"""Documents de sortie Word uniquement, sans conversion externe."""
import io
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt


def build_sortie_word(patient, kind, template=None):
    if kind not in ('crh', 'ordonnance'):
        raise ValueError('Type de document invalide.')
    template = Path(template) if template else Path(__file__).with_name('sortie_modele.docx')
    doc = Document(str(template))
    for child in list(doc._element.body):
        if not child.tag.endswith('}sectPr'):
            doc._element.body.remove(child)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    p = patient
    def date(value):
        value = str(value or '')
        return '/'.join(value.split('-')[::-1]) if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value) else value or 'À préciser'
    def add(text='', bold=False, center=False):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(5)
        if center:
            para.alignment = 1
        run = para.add_run(str(text or ''))
        run.bold = bold
        return para
    name = ' '.join(str(p.get(k) or '') for k in ('nom', 'prenom')).strip()
    if kind == 'crh':
        summary = str(p.get('sortie_resume') or '').strip()
        if not summary:
            raise ValueError('Renseignez et enregistrez le résumé avant le téléchargement du CRH.')
        add('Compte rendu d’hospitalisation', True, True)
        add('Nom et prénom : ' + name + ' — IP : ' + str(p.get('ip') or 'Non renseigné'), True)
        add('Âge : ' + (str(p['age']) + ' ans' if p.get('age') is not None else 'Non renseigné'))
        add('Date d’entrée : ' + date(p.get('date_entree')) + ' — Date de sortie : ' + date(p.get('sortie_date')))
        lines = summary.splitlines()
        # Anciens brouillons : ne pas répéter le titre et l'identité technique.
        for line in lines:
            stripped = line.strip()
            if stripped.upper() in ('COMPTE RENDU D’HOSPITALISATION', "COMPTE RENDU D'HOSPITALISATION"):
                continue
            if stripped.startswith('Nom et prénom :') or (stripped.startswith('Date d’entrée :') and 'Date de sortie :' in stripped):
                continue
            add(line, bool(stripped) and stripped.endswith(':'))
    else:
        prescription = str(p.get('sortie_ordonnance') or '').strip()
        if not prescription:
            raise ValueError('Aucun médicament enregistré dans l’ordonnance de sortie.')
        add('Oujda, le ' + date(p.get('sortie_date')))
        add(name, True)
        add('ORDONNANCE', True, True)
        number = 0
        for line in prescription.splitlines():
            if not line.strip():
                continue
            match = re.match(r'^• (.*?) — Posologie : (.*?) — Durée : (.*)$', line)
            if match:
                number += 1
                add(str(number) + '. ' + match[1], True)
                add(match[2] + (' — Durée : ' + match[3] if match[3] else ''))
                add()
            else:
                add(line)
    add()
    add('Signature et cachet', True)
    data = io.BytesIO()
    doc.save(data)
    data.seek(0)
    return data


def register_sortie_documents(app, db, login_required):
    from flask import jsonify, send_file
    @app.get('/api/patients/<int:pid>/sortie-document/<kind>/<fmt>')
    @login_required
    def sortie_document(pid, kind, fmt):
        if kind not in ('crh', 'ordonnance') or fmt != 'docx':
            return jsonify(error='Seul le téléchargement Word DOCX est disponible.'), 400
        row = db().execute('SELECT * FROM patients WHERE id=?', (pid,)).fetchone()
        if not row:
            return jsonify(error='Patient introuvable.'), 404
        patient = dict(row)
        if patient.get('status') != 'sorti':
            return jsonify(error='Validez la sortie avant de télécharger les documents.'), 409
        try:
            data = build_sortie_word(patient, kind)
        except ValueError as exc:
            return jsonify(error=str(exc)), 409
        except FileNotFoundError:
            return jsonify(error='Modèle sortie_modele.docx absent du serveur.'), 503
        label = re.sub(r'[^A-Za-z0-9_-]+', '_', str(patient.get('nom') or 'patient')).strip('_') or 'patient'
        response = send_file(data, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', as_attachment=True, download_name=kind + '_' + label + '_' + str(pid) + '.docx')
        response.headers['Cache-Control'] = 'no-store'
        return response
