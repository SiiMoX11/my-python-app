# STYLE FICHE TTT V2
import os
import json
import copy
import sqlite3
import datetime as dt
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR=os.path.dirname(os.path.abspath(__file__))

def catalogue():
    with open(os.path.join(BASE_DIR,'catalogue_traitements.json'),encoding='utf-8') as f:
        rows=json.load(f)
    with open(os.path.join(BASE_DIR,'ordre_traitements.json'),encoding='utf-8') as f:
        order=json.load(f)
    result=[{**row,'rang':order['ordre'].get(row['nom'],5)} for row in rows]
    return sorted(result,key=lambda row:row['rang'])


def treatment_rank(entry):
    mapping={row['nom']:row['rang'] for row in catalogue()}
    if entry.get('title') in mapping:return mapping[entry['title']]
    categories={'solutés':0,'solutes':0,'sérums':0,'serums':0,'antalgique':1,'antibiotique':2,'anticoagulant':3,'anticoagulation':3,'insuline':4}
    return categories.get(str(entry.get('category') or '').strip().lower(),5)


def treatment_data(e):
    try:p=json.loads(e.get('details') or '')
    except (ValueError,TypeError):p=None
    if isinstance(p,dict) and p.get('schema')=='ttt-catalogue-v1':return p
    return None

def set_patient_header(doc,patient,diseases):
    values={'Nom et Prénom':(patient.get('nom') or '')+' '+(patient.get('prenom') or ''),'IP':patient.get('ip') or '', 'Diagnostic':patient.get('motif') or '', 'Maladies associées':diseases or ''}
    replaced=set()
    for part in doc.part.package.parts:
        if 'header' not in str(part.partname):continue
        for paragraph in part.element.iter(qn('w:p')):
            if len(list(paragraph.iter(qn('w:p'))))!=1:continue
            nodes=list(paragraph.iter(qn('w:t')))
            text=''.join(n.text or '' for n in nodes).replace('\u00a0',' ').strip()
            for label,value in values.items():
                if text.replace('\u00a0',' ').startswith(label) and ':' in text:
                    if nodes:
                        nodes[0].text=label+' : '+str(value).strip()
                        nodes[0].set(qn('xml:space'),'preserve')
                        for node in nodes[1:]:node.text=''
                        replaced.add(label)
                    break
    if set(values)-replaced:raise ValueError('Les champs du modèle Word attendu sont introuvables.')

def set_cell(cell,text,bold=False,center=False):
    # Preserve the model's paragraph properties and first-run formatting.
    paragraph=cell.paragraphs[0]
    properties=None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        properties=copy.deepcopy(paragraph.runs[0]._r.rPr)
    elif paragraph._p.pPr is not None and paragraph._p.pPr.find(qn('w:rPr')) is not None:
        properties=copy.deepcopy(paragraph._p.pPr.find(qn('w:rPr')))
    for other in list(cell.paragraphs[1:]):
        cell._tc.remove(other._p)
    for child in list(paragraph._p):
        if child.tag!=qn('w:pPr'):
            paragraph._p.remove(child)
    run=paragraph.add_run(str(text or ''))
    if properties is not None:
        run._r.insert(0,properties)


def make_treatment_sheet(patient,diseases,treatments,stamp):
    start=dt.date.fromisoformat(stamp[:10])
    doc=Document(os.path.join(BASE_DIR,'modele_fiche_traitement.docx'))
    set_patient_header(doc,patient,diseases)
    if not doc.tables or len(doc.tables[0].columns)!=7:raise ValueError('Modèle de fiche de traitement incompatible.')
    table=doc.tables[0]
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    set_cell(table.cell(0,0),'Traitement prescrit',True,True)
    dates=[start+dt.timedelta(days=i) for i in range(6)]
    for i,day in enumerate(dates,1):set_cell(table.cell(0,i),day.strftime('%d/%m'),True,True)
    active=[e for e in treatments if not e.get('done')]
    active.sort(key=lambda e:(e.get('date') or '',e.get('id') or 0),reverse=True)
    active.sort(key=treatment_rank)
    while len(table.rows)<len(active)+1:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    available={row['nom']:row for row in catalogue()}
    for row in table.rows[1:]:
        for cell in row.cells:set_cell(cell,'')
    for i,e in enumerate(active,1):
        stored=treatment_data(e);preset=available.get(e.get('title'))
        # No inference of dose, frequency or duration from arbitrary text.
        takes=int(stored['prises']) if stored else 0
        planned=int(stored['jours']) if stored else 0
        title=e.get('title') or ''
        complement=(e.get('content') or '').strip()
        notes=(stored.get('notes') or '') if stored else (e.get('details') or '')
        text=title+('\n'+complement if complement else '')+('\n'+notes if notes else '')
        set_cell(table.cell(i,0),text)
        try:entry_start=dt.date.fromisoformat(e.get('date') or start.isoformat())
        except ValueError:entry_start=start
        plan_start=max(start,entry_start)
        for j,day in enumerate(dates,1):
            allowed=0<planned<=6 and plan_start<=day<plan_start+dt.timedelta(days=planned)
            set_cell(table.cell(i,j),'Δ'*takes if allowed and takes else '',True,True)
    buffer=BytesIO();doc.save(buffer);buffer.seek(0);return buffer

def register_treatments(app,db,login_required,now,log,db_path):
    from flask import request,jsonify,send_file
    conn=sqlite3.connect(db_path)
    try:
        conn.execute("CREATE TABLE IF NOT EXISTS patient_treatment_info(patient_id INTEGER PRIMARY KEY REFERENCES patients(id) ON DELETE CASCADE, associated_diseases TEXT NOT NULL DEFAULT '', updated_at TEXT NOT NULL)")
        conn.commit()
    finally:conn.close()

    @app.get('/api/ttt/catalogue')
    @login_required
    def api_ttt_catalogue():return jsonify(traitements=catalogue())

    @app.get('/api/patients/<int:pid>/ttt-profile')
    @login_required
    def api_ttt_profile(pid):
        if not db().execute('SELECT 1 FROM patients WHERE id=?',(pid,)).fetchone():return jsonify(error='Patient introuvable.'),404
        row=db().execute('SELECT associated_diseases FROM patient_treatment_info WHERE patient_id=?',(pid,)).fetchone()
        return jsonify(associated_diseases=row['associated_diseases'] if row else '')

    @app.put('/api/patients/<int:pid>/ttt-profile')
    @login_required
    def api_ttt_update_profile(pid):
        if not db().execute('SELECT 1 FROM patients WHERE id=?',(pid,)).fetchone():return jsonify(error='Patient introuvable.'),404
        data=request.get_json(silent=True) or {};value=data.get('associated_diseases','')
        if not isinstance(value,str) or len(value)>1000:return jsonify(error='Maladies associées : 1000 caractères maximum.'),400
        value=value.strip()
        db().execute('INSERT INTO patient_treatment_info(patient_id,associated_diseases,updated_at) VALUES(?,?,?) ON CONFLICT(patient_id) DO UPDATE SET associated_diseases=excluded.associated_diseases,updated_at=excluded.updated_at',(pid,value,now()))
        log(pid,'Maladies associées (TTT) modifiées',value);db().commit()
        return jsonify(ok=True)

    @app.before_request
    def validate_ttt_request():
        if request.endpoint not in {'api_add_entry','api_update_entry'}:return None
        @login_required
        def check():
            data=request.get_json(silent=True) or {};kind=data.get('kind');existing=None
            if request.endpoint=='api_update_entry':
                existing=db().execute('SELECT * FROM entries WHERE id=?',(request.view_args['eid'],)).fetchone()
                kind=existing['kind'] if existing else None
            if kind!='traitement':return None
            try:p=json.loads(data.get('details') or '')
            except (ValueError,TypeError):return None
            if not isinstance(p,dict) or p.get('schema')!='ttt-catalogue-v1':return None
            title=(data.get('title') or '').strip()
            names={r['nom'] for r in catalogue()}
            if title not in names and not (existing and title==existing['title']):return jsonify(error='Choisissez un traitement du catalogue.'),400
            try:
                takes=p.get('prises');days=p.get('jours')
                if type(takes) is not int or type(days) is not int:raise ValueError()
                dt.date.fromisoformat(data.get('date') or '')
                if not 0<=takes<=24 or not 1<=days<=6:raise ValueError()
            except (ValueError,TypeError):return jsonify(error='Vérifiez prises/jour (0–24), jours de fiche (1–6) et date.'),400
            notes=p.get('notes') or ''
            if not isinstance(notes,str) or len(notes)>4000:return jsonify(error='Remarques invalides.'),400
            data['title']=title
            data['details']=json.dumps({'schema':'ttt-catalogue-v1','prises':takes,'jours':days,'notes':notes.strip()},ensure_ascii=False)
            return None
        return check()

    @app.get('/api/patients/<int:pid>/fiche-traitement')
    @login_required
    def api_ttt_word(pid):
        conn=db();conn.execute('BEGIN')
        try:
            patient=conn.execute('SELECT * FROM patients WHERE id=?',(pid,)).fetchone()
            if not patient:return jsonify(error='Patient introuvable.'),404
            row=conn.execute('SELECT associated_diseases FROM patient_treatment_info WHERE patient_id=?',(pid,)).fetchone()
            treatments=[dict(e) for e in conn.execute("SELECT * FROM entries WHERE patient_id=? AND kind='traitement'",(pid,)).fetchall()]
            stamp=now();buffer=make_treatment_sheet(dict(patient),row['associated_diseases'] if row else '',treatments,stamp)
        except (ValueError,OSError) as exc:return jsonify(error=str(exc)),422
        finally:conn.rollback()
        response=send_file(buffer,as_attachment=True,download_name='Fiche_TTT_patient_'+str(pid)+'_'+stamp[:10]+'.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response.headers['Cache-Control']='no-store';return response
