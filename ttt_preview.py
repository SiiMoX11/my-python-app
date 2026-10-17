import base64
from html import escape
from docx import Document
from docx.oxml.ns import qn
from traitements_ttt import make_treatment_sheet


def paragraph_html(element):
    ppr=element.find(qn('w:pPr'))
    inherited=ppr.find(qn('w:rPr')) if ppr is not None else None
    def property_value(properties,tag):
        node=properties.find(qn('w:'+tag)) if properties is not None else None
        return node
    jc=property_value(ppr,'jc')
    align=jc.get(qn('w:val')) if jc is not None else 'left'
    align=align if align in {'left','right','center','justify'} else 'left'
    chunks=[]
    for run in element:
        if run.tag!=qn('w:r'):continue
        rpr=run.find(qn('w:rPr'))
        size=property_value(rpr,'sz')
        if size is None:size=property_value(inherited,'sz')
        points=12
        if size is not None:
            try:points=float(size.get(qn('w:val')))/2
            except (ValueError,TypeError):pass
        points=max(6,min(72,points))
        bold=property_value(rpr,'b')
        if bold is None:bold=property_value(inherited,'b')
        weight='bold' if bold is not None and bold.get(qn('w:val'),'1') not in {'0','false','off'} else 'normal'
        text=[]
        for node in run.iter():
            if node.tag==qn('w:t'):text.append(escape(node.text or ''))
            elif node.tag==qn('w:br'):text.append('<br>')
        if text:chunks.append('<span style="font-size:'+str(points)+'pt;font-weight:'+weight+'">'+''.join(text)+'</span>')
    return '<p dir="auto" style="text-align:'+align+'">'+''.join(chunks)+'</p>'


def sheet_html(buffer):
    buffer.seek(0);doc=Document(buffer);header=[];seen=set();logos=[];seen_images=set()
    for part in doc.part.package.parts:
        if 'header' not in str(part.partname):continue
        for p in part.element.iter(qn('w:p')):
            if len(list(p.iter(qn('w:p'))))!=1:continue
            text=''.join(n.text or '' for n in p.iter(qn('w:t'))).strip()
            if text and text not in seen:
                seen.add(text);header.append(paragraph_html(p))
        for rel in part.rels.values():
            if rel.is_external or not rel.reltype.endswith('/image'):continue
            image=rel.target_part
            if image.partname in seen_images:continue
            seen_images.add(image.partname)
            if image.content_type not in {'image/png','image/jpeg','image/gif'}:continue
            logos.append('<img alt="Logo du modèle" src="data:'+image.content_type+';base64,'+base64.b64encode(image.blob).decode('ascii')+'">')
    tables=[]
    for table in doc.tables:
        rows=[]
        for row in table.rows:
            cells=['<td>'+''.join(paragraph_html(p._p) for p in cell.paragraphs)+'</td>' for cell in row.cells]
            rows.append('<tr>'+''.join(cells)+'</tr>')
        tables.append('<table>'+''.join(rows)+'</table>')
    foot=[];seen_footer=set()
    for part in doc.part.package.parts:
        if 'footer' not in str(part.partname):continue
        for p in part.element.iter(qn('w:p')):
            if len(list(p.iter(qn('w:p'))))!=1:continue
            text=''.join(n.text or '' for n in p.iter(qn('w:t'))).strip()
            if text and text not in seen_footer:
                seen_footer.add(text);foot.append(paragraph_html(p))
    return '<!doctype html><html lang="fr"><head><meta charset="utf-8"><style>body{font-family:Cambria,"Times New Roman",serif;background:#e2e8f0;margin:0;color:#000}.sheet{background:white;max-width:820px;margin:12px auto;padding:24px;box-sizing:border-box}p{margin:5px 0;white-space:pre-wrap;line-height:1.3}.logos{display:flex;justify-content:space-between;gap:15px}.logos img{max-height:80px;max-width:150px;object-fit:contain}table{border-collapse:collapse;width:100%;margin-top:16px;table-layout:fixed}td{border:1px solid #000;padding:8px 4px;text-align:center;vertical-align:middle;overflow-wrap:anywhere}td:first-child{width:30%}footer{font-size:10pt;margin-top:16px}footer p span{font-size:10pt!important}@media(max-width:600px){.sheet{min-width:680px;margin:0;padding:15px}body{overflow:auto}}</style></head><body><div class="sheet"><header><div class="logos">'+''.join(logos)+'</div>'+''.join(header)+'</header>'+''.join(tables)+'<footer>'+''.join(foot)+'</footer></div></body></html>'


def register_ttt_preview(app,db,login_required,now):
    from flask import jsonify
    @app.get('/api/patients/<int:pid>/fiche-traitement/preview')
    @login_required
    def api_ttt_preview(pid):
        conn=db();conn.execute('BEGIN')
        try:
            patient=conn.execute('SELECT * FROM patients WHERE id=?',(pid,)).fetchone()
            if not patient:return jsonify(error='Patient introuvable.'),404
            row=conn.execute('SELECT associated_diseases FROM patient_treatment_info WHERE patient_id=?',(pid,)).fetchone()
            treatments=[dict(e) for e in conn.execute("SELECT * FROM entries WHERE patient_id=? AND kind='traitement'",(pid,)).fetchall()]
            stamp=now();buffer=make_treatment_sheet(dict(patient),row['associated_diseases'] if row else '',treatments,stamp)
            original=buffer.getvalue();html=sheet_html(buffer)
        except (ValueError,OSError) as exc:return jsonify(error=str(exc)),422
        finally:conn.rollback()
        response=jsonify(html=html,docx_base64=base64.b64encode(original).decode('ascii'),filename='Fiche_TTT_'+str(pid)+'_'+stamp[:10]+'.docx')
        response.headers['Cache-Control']='no-store';return response
