import json
import re
import math
import datetime as dt
from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOM_ORDER = [str(n) for n in range(210, 219)] + [str(n) for n in range(201, 210)]
GROUPS = [('Alb',['Alb']),('GB',['GB']),('Hb',['Hb']),('Plq',['Plq']),('TP',['TP']),('CRP',['CRP']),('Lipase',['Lipase']),('BT/BD',['BT','BD']),('GGT/PAL',['GGT','PAL']),('ASAT/ALAT',['ASAT','ALAT']),('Urée/Créat',['Uree','Creat']),('Na/K',['Na','K']),('ECBU',['ECBU']),('ACE',['ACE']),('CA19-9',['CA19-9']),('CA125',['CA125']),('CA15-3',['CA15-3']),('Autre',['Autre'])]
BLACK = RGBColor(0,0,0)
RED = RGBColor(200,30,30)

def bed_key(p):
    return ROOM_ORDER.index(str(p['room'])), 0 if p['position']=='P' else 1, p['id']

def date_fr(value):
    value = str(value or '')
    match = re.match(r'^(\d{4})-(\d{2})-(\d{2})',value)
    return '/'.join((match[3],match[2],match[1])) if match else value

def ordered(items, reverse=False):
    return sorted(items, key=lambda e:(e.get('date') or '', e.get('id') or 0), reverse=reverse)

def paragraph(cell, text='', bold=False, center=False, underline=False, highlight=None):
    p=cell.add_paragraph()
    p.paragraph_format.space_after=Pt(4)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(str(text or ''))
    run.bold=bold;run.underline=underline;run.font.color.rgb=BLACK
    if highlight:run.font.highlight_color=highlight
    return p

def body_text(cell,e):
    if e.get('content'):paragraph(cell,e['content'])
    if e.get('details'):paragraph(cell,e['details'])

def entry(cell,e,label=''):
    heading=' '.join(str(v) for v in (label,e.get('category'),e.get('title'),date_fr(e.get('date'))) if v)
    if heading:paragraph(cell,heading+' :',bold=True,underline=True)
    body_text(cell,e)

def numeric(value,key):
    text=str(value if value is not None else '').strip().replace(',','.')
    if key in {'GB','Plq'} and re.fullmatch(r'\d+(?:\.\d+)?[kK]',text):
        return float(text[:-1])*1000
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?',text):return None
    n=float(text)
    return n if math.isfinite(n) else None

def decode_bilan(e):
    try:p=json.loads(e.get('details') or '')
    except (ValueError,TypeError):return None
    if not isinstance(p,dict) or not isinstance(p.get('values'),dict):return None
    rows={}
    if p.get('schema')=='bilan-simple-v2':
        for key,value in p['values'].items():
            if str(value if value is not None else '').strip():
                rows[key]={'text':str(value),'reference':p.get('references',{}).get(key,{})}
    elif p.get('schema')=='bilan-v1':
        for key,row in p['values'].items():
            if isinstance(row,dict) and str(row.get('value') or '').strip():
                lo=numeric(row.get('min'),key);hi=numeric(row.get('max'),key)
                rows[key]={'text':str(row['value']),'reference':{'active':bool(row.get('unit')),'min':lo,'max':hi}}
    else:return None
    return rows

def abnormal(row,key):
    r=row.get('reference') or {};n=numeric(row.get('text'),key)
    if n is None or not r.get('active'):return False
    lo=numeric(r.get('min'),key);hi=numeric(r.get('max'),key)
    return (lo is not None and n<lo) or (hi is not None and (n>=hi if r.get('max_exclusive') else n>hi))

def short_dates(days):
    last=days[-1] if days else ''
    labels=[]
    for i,day in enumerate(days):
        if i==len(days)-1 or len(day)!=10:labels.append(date_fr(day))
        elif day[:7]==last[:7]:labels.append(day[8:10])
        elif day[:4]==last[:4]:labels.append(day[8:10]+'/'+day[5:7])
        else:labels.append(date_fr(day))
    return labels

def bilan_kinetic(cell,items):
    groups={};other=[];legacy=[]
    for e in ordered(items):
        rows=decode_bilan(e)
        if rows is None:
            legacy.append(e);continue
        day=e.get('date') or ''
        if rows.get('Autre'):other.append((day,e.get('id') or 0,rows['Autre']['text']))
        values={key:row for key,row in rows.items() if key!='Autre'}
        if values:groups.setdefault(day,{}).update(values)
    days=sorted(groups)
    if days:
        heading=paragraph(cell,'BB : ',bold=True,underline=True)
        labels=short_dates(days)
        for i,label in enumerate(labels):
            if i:heading.add_run(' → ').font.color.rgb=BLACK
            run=heading.add_run(label);run.bold=True;run.font.color.rgb=BLACK
            if i==len(days)-1:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
        for label,keys in GROUPS:
            if 'Autre' in keys or not any(any(k in groups[day] for k in keys) for day in days):continue
            p=cell.add_paragraph()
            p.paragraph_format.space_after=Pt(3)
            p.add_run(label+' : ').bold=True
            for i,day in enumerate(days):
                if i:p.add_run(' → ').font.color.rgb=BLACK
                for k,key in enumerate(keys):
                    if k:p.add_run('/').font.color.rgb=BLACK
                    row=groups[day].get(key)
                    run=p.add_run(row['text'] if row else '—')
                    latest=i==len(days)-1 and row is not None
                    run.bold=latest
                    run.font.color.rgb=RED if latest and abnormal(row,key) else BLACK
                    if latest:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
    # Other results remain outside kinetics, newest first.
    for day,ident,text in sorted(other,key=lambda x:(x[0],x[1]),reverse=True):
        paragraph(cell,'Autre — '+date_fr(day),bold=True)
        paragraph(cell,text)
    for e in ordered(legacy,reverse=True):
        entry(cell,e,'BB')

def done(value):
    return str(value).strip().lower() in {'1','true','yes'}

def post_label(e,stamp):
    text=str(e.get('content') or '')
    if e.get('category')!='Gestes faits':return text
    try:
        n=(dt.date.fromisoformat(stamp[:10])-dt.date.fromisoformat(e.get('date') or '')).days
    except ValueError:return text
    if n<0:return text
    text=re.sub(r'^\s*J\s*\d+\s*(?:post\b\s*)?','',text,flags=re.I)
    text=re.sub(r'^\s*post\b\s*','',text,flags=re.I).strip()
    return 'J'+str(n)+' post '+text

def clinical(cell,items,stamp):
    gestures=ordered([e for e in items if e.get('category')=='Gestes faits'],True)
    drains=ordered([e for e in items if e.get('category')=='Drains'],True)
    rest=ordered([e for e in items if e.get('category') not in {'Gestes faits','Drains'}],True)
    if gestures:
        paragraph(cell,'Geste :',bold=True,underline=True)
        for e in gestures:
            paragraph(cell,post_label(e,stamp))
            if e.get('details'):paragraph(cell,e['details'])
    if drains:
        paragraph(cell,'Drains :',bold=True,underline=True)
        for e in drains:
            body_text(cell,e)
    names={'Surveillance clinique':'Surv.','Constantes':'Constantes'}
    for e in rest:
        label=names.get(e.get('category'),e.get('category') or 'Clinique')
        paragraph(cell,label+' :',bold=True,underline=True)
        body_text(cell,e)

def build_document(patients,entries,stamp):
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(1)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(3);normal.paragraph_format.line_spacing=1
    doc.add_paragraph().add_run(date_fr(stamp[:10])+' à '+stamp[11:16]+' — heure plateforme').bold=True
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('SERVICE CHIRURGIE VISCÉRALE A');r.bold=True;r.font.size=Pt(13)
    table=doc.add_table(rows=0,cols=5);table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[1.45,3.7,3.2,6.1,4.55]
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    for patient in sorted(patients,key=bed_key):
        cells=table.add_row().cells
        for cell,width in zip(cells,widths):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell._element.remove(cell.paragraphs[0]._p)
        paragraph(cells[0],str(patient['room'])+' '+patient['position'],True,True)
        paragraph(cells[1],patient['nom']+' '+(patient.get('prenom') or ''),True,True)
        if patient.get('age') is not None:paragraph(cells[1],str(patient['age'])+' ans',True,True)
        paragraph(cells[1],'IP : '+str(patient.get('ip') or '—'),True,True)
        paragraph(cells[1],patient.get('mutuelle') or '—',True,True)
        if patient.get('groupe'):paragraph(cells[1],patient['groupe'],True,True,highlight=WD_COLOR_INDEX.TURQUOISE)
        paragraph(cells[2],patient.get('motif') or 'Motif non renseigné',True,True)
        if patient.get('date_entree'):paragraph(cells[2],'Entrée : '+date_fr(patient['date_entree']),center=True)
        E=entries.get(patient['id'],{})
        paragraph(cells[2],'ATCDs :',True,True)
        for e in E.get('atcd',[]):body_text(cells[2],e)
        if not E.get('atcd'):paragraph(cells[2],'Non renseignés',center=True)
        bilan_kinetic(cells[3],E.get('bilan',[]))
        for e in ordered(E.get('examen',[]),True):entry(cells[3],e)
        pending=[e for e in E.get('consigne',[]) if not done(e.get('done',0))]
        if pending:
            paragraph(cells[4],'Consignes à faire :',True,underline=True)
            for e in pending:paragraph(cells[4],e.get('content') or e.get('title') or '',True,highlight=WD_COLOR_INDEX.YELLOW)
        clinical(cells[4],E.get('clinique',[]),stamp)
        if E.get('avis'):
            paragraph(cells[4],'Avis :',True,underline=True)
            for e in ordered(E['avis'],True):entry(cells[4],e,'Avis')
        for cell in cells:
            if not cell.paragraphs:cell.add_paragraph('—')
    if not patients:doc.add_paragraph('Aucun patient hospitalisé dans les chambres du service.')
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Document confidentiel — Page ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    buffer=BytesIO();doc.save(buffer);buffer.seek(0);return buffer

def register_tableau(app, db, login_required, now):
    from flask import send_file, jsonify
    @app.get("/api/tableau/word")
    @login_required
    def api_tableau_word():
        conn = db()
        conn.execute("BEGIN")
        try:
            patients = [dict(r) for r in conn.execute("SELECT * FROM patients WHERE status='hospitalise' AND room IS NOT NULL AND position IN ('P','F') AND (hors_service IS NULL OR hors_service='')").fetchall() if str(r["room"]) in ROOM_ORDER]
            entries = {}
            if patients:
                ids = [p["id"] for p in patients]
                marks = ",".join("?" for _ in ids)
                rows = conn.execute(f"SELECT * FROM entries WHERE patient_id IN ({marks}) AND kind IN ('atcd','clinique','bilan','examen','avis','consigne') ORDER BY date DESC,id DESC", ids).fetchall()
                for r in rows:
                    e = dict(r)
                    entries.setdefault(e["patient_id"], {}).setdefault(e["kind"], []).append(e)
            stamp = now()
        finally:
            conn.rollback()
        buffer = build_document(patients, entries, stamp)
        response = send_file(buffer, as_attachment=True, download_name="Tableau_service_" + stamp[:10] + ".docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response.headers["Cache-Control"] = "no-store"
        return response
